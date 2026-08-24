#!/usr/bin/env python
from __future__ import annotations

from pathlib import Path

from legend_text import ALL_TABLE_FIGURE_LEGENDS, MAIN_TABLE_FIGURE_LEGENDS, SUPPLEMENTARY_TABLE_FIGURE_LEGENDS
from utils import MANUSCRIPT_DIR, REPORT_DIR, SUPPLEMENT_DIR, ensure_dirs


def write_docx_from_markdown_text(text: str, path: Path) -> bool:
    try:
        from docx import Document
    except Exception:
        return False

    doc = Document()
    for line in text.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.strip() == "":
            continue
        else:
            doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return True


def main() -> None:
    ensure_dirs()

    outputs = [
        (MANUSCRIPT_DIR / "table_figure_legends.md", MAIN_TABLE_FIGURE_LEGENDS),
        (SUPPLEMENT_DIR / "supplementary_table_figure_legends.md", SUPPLEMENTARY_TABLE_FIGURE_LEGENDS),
        (MANUSCRIPT_DIR / "all_table_figure_legends.md", ALL_TABLE_FIGURE_LEGENDS),
    ]

    docx_status = []
    for md_path, text in outputs:
        md_path.write_text(text.rstrip() + "\n", encoding="utf-8")
        docx_path = md_path.with_suffix(".docx")
        docx_status.append((docx_path, write_docx_from_markdown_text(text, docx_path)))

    report_lines = [
        "Table and figure legends report",
        "===============================",
        "",
        "Generated files:",
    ]
    report_lines.extend(f"- {path}" for path, _ in outputs)
    report_lines.extend(f"- {path} ({'written' if ok else 'python-docx unavailable'})" for path, ok in docx_status)
    report_lines.extend(
        [
            "",
            "Coverage:",
            "- Main tables: Table 1, Table 2, Table 3",
            "- Main figures: Figure 1, Figure 2, Figure 3, Figure 4, Figure 5",
            "- Supplementary tables: Table S1, Table S2, Table S3, Table S4, Table S5, Table S6, Table S7",
            "- Supplementary figures: Supplementary Figure S1, S2, S3, S4",
            "",
        ]
    )
    (REPORT_DIR / "table_figure_legends_report.txt").write_text("\n".join(report_lines), encoding="utf-8")
    print("Table and figure legends written.")
    for path, _ in outputs:
        print(path)
    for path, ok in docx_status:
        print(f"{path} ({'written' if ok else 'python-docx unavailable'})")


if __name__ == "__main__":
    main()
