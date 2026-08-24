"""Build the revised main tables as tables.docx."""
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

sys.path.insert(0, str(Path(__file__).resolve().parent))
import paths as P

P.ensure_dirs()
OUT = P.RESULTS_DIR
DEST = P.TABLE_DIR / "tables.docx"

doc = Document()
for s in doc.sections:
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Inches(11.69), Inches(8.27)
    s.left_margin = s.right_margin = Inches(0.6)
    s.top_margin = s.bottom_margin = Inches(0.7)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(9)


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def add_table(title, df, note, bold_rows=()):
    h = doc.add_paragraph()
    r = h.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    h.paragraph_format.space_after = Pt(6)

    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, col in enumerate(df.columns):
        c = t.rows[0].cells[j]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(col))
        run.bold = True
        run.font.size = Pt(8.5)
        shade(c, "D9E2F3")
    for i in range(len(df)):
        cells = t.add_row().cells
        for j, col in enumerate(df.columns):
            raw = df.iloc[i, j]
            v = "" if pd.isna(raw) else str(raw)
            p = cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(v)
            run.font.size = Pt(8.5)
            if i in bold_rows:
                run.bold = True
            if str(col).strip() == "Instability flag" and v:
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    n = doc.add_paragraph()
    nr = n.add_run(note)
    nr.italic = True
    nr.font.size = Pt(8)
    n.paragraph_format.space_before = Pt(6)
    n.paragraph_format.space_after = Pt(14)


def money(x):
    return f"{int(round(float(x))):,}" if pd.notna(x) and str(x) != "" else ""


def ci(lo, hi):
    """Range string; uses 'to' when a bound is negative so the minus sign stays readable."""
    if pd.isna(lo) or pd.isna(hi):
        return ""
    sep = " to " if float(lo) < 0 else "-"
    return f"{money(lo)}{sep}{money(hi)}"


RENAME = {"other selected ball sports": "other selected ball and stick sports"}

# ------------------------------------------------------------------- Table 1
t1 = pd.read_csv(OUT / "table1_with_ci.csv")
t1["label"] = t1["label"].replace(RENAME)
T1 = pd.DataFrame({
    "Characteristic": t1.apply(lambda r: f"{r.section}: {r.label}", axis=1),
    "Unweighted n": t1["unweighted_n"],
    "Weighted national estimate": t1["weighted_estimate"].map(money),
    "95% CI": t1.apply(lambda r: ci(r.ci_lo, r.ci_hi), axis=1),
    "CV": t1["cv_pct"].map(lambda v: f"{v:.1f}%"),
    "Weighted %": t1["weighted_pct"].map(lambda v: f"{v:.1f}%"),
    "95% CI for %": t1.apply(lambda r: f"{r.pct_ci_lo:.1f}-{r.pct_ci_hi:.1f}", axis=1),
    "Instability flag": t1["instability_flag"].fillna(""),
})
add_table(
    "Table 1. Baseline characteristics of the main analytic cohort, 2015-2025",
    T1,
    "Weighted national estimates of emergency department visits with 95% confidence intervals. Confidence intervals for "
    "totals were computed as the estimate ± 1.96 standard errors, with standard errors estimated by Taylor linearization "
    "using NEISS stratum, primary sampling unit, and sample-weight variables; intervals for percentages were computed on "
    "the log scale from the linearized ratio variance. Instability flags follow CPSC criteria (weighted estimate <1,200, "
    "unweighted n <20, or CV >33%). All weighted estimates are rounded from unrounded survey-weighted values, so components "
    "may not sum exactly to the cohort total of 441,441 (exact value 441,440.82). Sex reflects emergency department "
    "administrative coding; the single record coded unknown/other is retained here but was excluded from the model in "
    "Table 3. CI, confidence interval; CV, coefficient of variation; NEISS, National Electronic Injury Surveillance System.",
)
doc.add_page_break()

# ------------------------------------------------------------------- Table 2
t2 = pd.read_csv(OUT / "table2_with_ci_and_severe_flags.csv")
t2["sport_group"] = t2["sport_group"].replace(RENAME)
T2 = pd.DataFrame({
    "a": t2["sport_group"],
    "b": t2["unweighted_n"],
    "c": t2["weighted_estimate"].map(money),
    "d": t2["est_ci"],
    "e": t2["cv_pct"].map(lambda v: f"{v:.1f}%"),
    "f": t2["fracture_pct"].map(lambda v: f"{v:.1f}%"),
    "g": t2["fracture_ci"],
    "h": t2["dislocation_pct"].map(lambda v: f"{v:.1f}%"),
    "i": t2["dislocation_ci"],
    "j": t2["top_season"],
})
T2.columns = ["Sport group", "Unweighted n", "Weighted national estimate", "95% CI",
              "CV of sport total", "Fracture %", "95% CI", "Dislocation %", "95% CI", "Top season"]
add_table(
    "Table 2. Sport-specific injury patterns, 2015-2025",
    T2,
    "The CV column refers to the coefficient of variation of each sport group's TOTAL weighted estimate. Sport-specific "
    "severe-disposition percentages are not reported: all six sport-specific severe-disposition estimates failed CPSC "
    "instability criteria, none reaching the reportability threshold of 1,200 (Supplementary Table S7). The overall "
    "severe-disposition estimate was 2,102 visits (95% CI 1,289-2,916; 0.48%, 95% CI 0.32-0.70). Fracture and dislocation "
    "percentages are weighted within sport group and sum to 100% for each sport. \"Other selected ball and stick sports\" "
    "combines verified official NEISS product/activity codes; its full weighted composition is given in Supplementary "
    "Table S1. CI, confidence interval; CV, coefficient of variation; NEISS, National Electronic Injury Surveillance System.",
)
doc.add_page_break()

# ------------------------------------------------------------------- Table 3
T3 = pd.DataFrame([
    ["Sport group: baseball/softball vs football", "0.66", "0.51-0.86", "0.002", "0.008"],
    ["Sport group: basketball vs football", "1.05", "0.92-1.20", "0.437", "0.486"],
    ["Sport group: other selected ball and stick sports vs football", "0.57", "0.41-0.80", "0.001", "0.005"],
    ["Sport group: soccer vs football", "0.82", "0.59-1.13", "0.217", "0.486"],
    ["Sport group: volleyball vs football", "1.26", "0.91-1.75", "0.162", "0.486"],
    ["Age group: 13-15 vs 10-12 years", "1.84", "1.59-2.14", "<0.001", "not applicable"],
    ["Age group: 16-18 vs 10-12 years", "4.86", "3.94-5.99", "<0.001", "not applicable"],
    ["Sex: records coded female vs records coded male", "0.57", "0.47-0.70", "<0.001", "not applicable"],
    ["Treatment year, per year", "0.99", "0.96-1.01", "0.357", "not applicable"],
], columns=["Predictor", "Adjusted OR", "95% CI", "p value (unadjusted)", "p value (Holm-adjusted)"])
add_table(
    "Table 3. Survey-weighted adjusted model for dislocation versus fracture, 2015-2025",
    T3,
    "THE ESTIMAND IS CONDITIONAL: the model estimates the odds that a presenting injury was coded as a dislocation rather "
    "than a fracture, among adolescents who presented to an emergency department with one of these two diagnoses. It is a "
    "comparison of diagnostic case mix among presenting patients and is not an estimate of injury risk. Estimated using "
    "survey::svyglm (binomial family, logit link) with NEISS stratum, primary sampling unit, and sample-weight variables. "
    "Reference categories: football, ages 10-12 years, records coded as male, and fracture diagnosis. Treatment year was "
    "modeled as a single continuous mean-centered term. Holm adjustment was applied to the family of five sport comparisons "
    "against football; age, sex, and year were prespecified adjustment variables and were not included in the multiplicity "
    "family. Sport-specific comparisons should be considered exploratory. Analytic n = 16,065; the single record coded "
    "unknown/other for sex was excluded to avoid sparse-cell separation. CI, confidence interval; OR, odds ratio.",
)

DEST.parent.mkdir(parents=True, exist_ok=True)
doc.save(DEST)
print("saved", DEST)
