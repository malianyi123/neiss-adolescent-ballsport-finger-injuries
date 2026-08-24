from __future__ import annotations

import csv
import math
import os
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
import yaml


ROOT = Path(__file__).resolve().parents[1]
CONFIG_DIR = ROOT / "config"
RAW_DIR = ROOT / "data" / "raw" / "neiss"
INTERMEDIATE_DIR = ROOT / "data" / "intermediate"
PROCESSED_DIR = ROOT / "data" / "processed"
OUTPUT_DIR = ROOT / "outputs"
REPORT_DIR = OUTPUT_DIR / "reports"
TABLE_DIR = OUTPUT_DIR / "tables"
FIGURE_DIR = OUTPUT_DIR / "figures"
SUPPLEMENT_DIR = OUTPUT_DIR / "supplementary"
MANUSCRIPT_DIR = OUTPUT_DIR / "manuscript"
DOCS_DIR = ROOT / "docs"


def ensure_dirs() -> None:
    for p in [
        CONFIG_DIR,
        RAW_DIR,
        INTERMEDIATE_DIR,
        PROCESSED_DIR,
        REPORT_DIR,
        TABLE_DIR,
        FIGURE_DIR,
        SUPPLEMENT_DIR,
        SUPPLEMENT_DIR / "supplementary_figures",
        MANUSCRIPT_DIR,
        DOCS_DIR,
    ]:
        p.mkdir(parents=True, exist_ok=True)


def load_yaml(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def write_yaml(obj: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        yaml.safe_dump(obj, f, sort_keys=False, allow_unicode=False)


def load_config() -> dict[str, Any]:
    return load_yaml(CONFIG_DIR / "study_config.yml")


def load_mappings() -> dict[str, Any]:
    path = CONFIG_DIR / "code_mappings_reviewed.yml"
    obj = load_yaml(path)
    if not obj.get("validated"):
        raise RuntimeError(
            "Official code mappings are not validated. Run scripts/01_parse_dictionary.py first."
        )
    return obj


def now_utc() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def annual_raw_path(year: int) -> Path:
    return RAW_DIR / f"neiss{year}.tsv"


def clean_year_path(year: int) -> Path:
    return INTERMEDIATE_DIR / f"neiss_clean_{year}.csv"


def run(cmd: list[str], timeout: int | None = None, quiet: bool = False) -> subprocess.CompletedProcess:
    if not quiet:
        print("+ " + " ".join(cmd), flush=True)
    return subprocess.run(cmd, check=False, capture_output=True, text=True, timeout=timeout)


def curl_download(url: str, dest: Path, timeout: int = 300) -> tuple[bool, str]:
    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp = dest.with_suffix(dest.suffix + ".part")
    headers = dest.with_suffix(dest.suffix + ".headers")
    cmd = [
        "curl",
        "-L",
        "--fail",
        "--retry",
        "3",
        "--retry-delay",
        "2",
        "--connect-timeout",
        "30",
        "-s",
        "-D",
        str(headers),
        "-o",
        str(tmp),
        url,
    ]
    proc = run(cmd, timeout=timeout, quiet=True)
    if proc.returncode != 0:
        if tmp.exists():
            tmp.unlink()
        return False, proc.stderr.strip() or proc.stdout.strip() or f"curl exited {proc.returncode}"
    if tmp.stat().st_size < 100:
        msg = f"downloaded file too small: {tmp.stat().st_size} bytes"
        tmp.unlink(missing_ok=True)
        return False, msg
    tmp.replace(dest)
    return True, headers.read_text(errors="replace") if headers.exists() else ""


def normalize_col(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


CANONICAL_COLUMN_ALIASES = {
    "cpsc_case_number": "case_id",
    "case_number": "case_id",
    "case_no": "case_id",
    "nek": "case_id",
    "treatment_date": "treatment_date",
    "trmt_date": "treatment_date",
    "date": "treatment_date",
    "age": "age",
    "sex": "sex",
    "race": "race",
    "other_race": "other_race",
    "hispanic": "hispanic",
    "ethnicity": "ethnicity",
    "body_part": "body_part",
    "bdpt": "body_part",
    "body_part_2": "body_part_2",
    "bdpt2": "body_part_2",
    "diagnosis": "diagnosis",
    "diag": "diagnosis",
    "diagnosis_2": "diagnosis_2",
    "diag2": "diagnosis_2",
    "other_diagnosis": "other_diagnosis",
    "other_diagnosis_2": "other_diagnosis_2",
    "disposition": "disposition",
    "disp": "disposition",
    "location": "location",
    "locale": "location",
    "fire_involvement": "fire_involvement",
    "product_1": "product_1",
    "prod1": "product_1",
    "product_2": "product_2",
    "prod2": "product_2",
    "product_3": "product_3",
    "prod3": "product_3",
    "alcohol": "alcohol",
    "alcohol_involved": "alcohol",
    "drug": "drug",
    "drug_involved": "drug",
    "narrative": "narrative",
    "narrative_1": "narrative",
    "narr1": "narrative",
    "stratum": "stratum",
    "psu": "psu",
    "weight": "weight",
}


REQUIRED_STANDARD_COLUMNS = [
    "case_id",
    "treatment_date",
    "year",
    "age",
    "sex",
    "race",
    "hispanic",
    "body_part",
    "diagnosis",
    "body_part_2",
    "diagnosis_2",
    "disposition",
    "location",
    "product_1",
    "product_2",
    "product_3",
    "narrative",
    "weight",
    "stratum",
    "psu",
]


def standardize_columns(df: pd.DataFrame) -> pd.DataFrame:
    renamed = {}
    for c in df.columns:
        norm = normalize_col(c)
        renamed[c] = CANONICAL_COLUMN_ALIASES.get(norm, norm)
    out = df.rename(columns=renamed)
    for c in REQUIRED_STANDARD_COLUMNS:
        if c not in out.columns:
            out[c] = pd.NA
    return out


def read_raw_neiss(path: Path) -> pd.DataFrame:
    # Older annual NEISS TSVs can contain extended bytes in narrative text.
    # Raw files remain unchanged; latin-1 is a lossless single-byte read path.
    return pd.read_csv(path, sep="\t", dtype=str, keep_default_na=False, na_values=[""], encoding="latin1")


def to_numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce")


def clean_numeric_code(series: pd.Series) -> pd.Series:
    num = pd.to_numeric(series.replace({"": np.nan, " ": np.nan}), errors="coerce")
    non_integer = num.notna() & ~np.isclose(num, np.round(num))
    num = num.mask(non_integer)
    return num.round().astype("Int64")


def parse_treatment_date(series: pd.Series) -> pd.Series:
    return pd.to_datetime(series, errors="coerce", format="%m/%d/%Y")


def load_clean_years(years: Iterable[int]) -> pd.DataFrame:
    frames = []
    missing = []
    for y in years:
        p = clean_year_path(int(y))
        if p.exists():
            frames.append(pd.read_csv(p, dtype={"case_id": str}))
        else:
            missing.append(str(p))
    if missing:
        raise FileNotFoundError("Missing cleaned yearly files:\n" + "\n".join(missing))
    return pd.concat(frames, ignore_index=True)


def label_from_mapping(mapping: dict[str, Any], code_family: str, code: int) -> str:
    try:
        return mapping[code_family][str(int(code))]["label"]
    except Exception:
        return ""


def sport_code_to_group(mappings: dict[str, Any]) -> dict[int, str]:
    out: dict[int, str] = {}
    for group, meta in mappings["sport_groups"].items():
        label = meta.get("label", group.replace("_", " "))
        for item in meta["codes"]:
            out[int(item["code"])] = label
    return out


def format_int(x: float | int | None) -> str:
    if x is None or pd.isna(x):
        return ""
    return f"{float(x):,.0f}"


def format_pct(x: float | int | None) -> str:
    if x is None or pd.isna(x):
        return ""
    return f"{float(x) * 100:.1f}%"


def format_or(x: float | int | None) -> str:
    if x is None or pd.isna(x):
        return ""
    return f"{float(x):.2f}"


def md_table(df: pd.DataFrame) -> str:
    if df.empty:
        return "_No rows generated._"
    out = df.copy()
    out = out.astype(object).where(pd.notna(out), "")
    headers = [str(c) for c in out.columns]

    def cell(x: Any) -> str:
        text = str(x)
        text = text.replace("|", "\\|").replace("\n", " ")
        return text

    lines = [
        "| " + " | ".join(cell(h) for h in headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for _, row in out.iterrows():
        lines.append("| " + " | ".join(cell(v) for v in row.tolist()) + " |")
    return "\n".join(lines)


def write_markdown_table(df: pd.DataFrame, path: Path, title: str | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    text = ""
    if title:
        text += f"# {title}\n\n"
    text += md_table(df) + "\n"
    path.write_text(text, encoding="utf-8")


def write_docx_tables(tables: list[tuple[str, pd.DataFrame]], path: Path) -> bool:
    try:
        from docx import Document
    except Exception:
        return False

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for idx, (title, df) in enumerate(tables):
        if idx:
            doc.add_page_break()
        doc.add_heading(title, level=1)
        if df.empty:
            doc.add_paragraph("No rows generated.")
            continue
        table = doc.add_table(rows=1, cols=len(df.columns))
        table.style = "Table Grid"
        for j, col in enumerate(df.columns):
            table.rows[0].cells[j].text = str(col)
        for _, row in df.iterrows():
            cells = table.add_row().cells
            for j, val in enumerate(row):
                cells[j].text = "" if pd.isna(val) else str(val)
    doc.save(path)
    return True


def instability_flag(
    n: int,
    estimate: float,
    cv: float | None,
    config: dict[str, Any],
) -> str:
    rules = config["instability"]
    flags = []
    if pd.isna(estimate) or estimate < float(rules["min_estimate"]):
        flags.append("estimate<1200")
    if int(n) < int(rules["min_unweighted_n"]):
        flags.append("n<20")
    if cv is not None and not pd.isna(cv) and cv > float(rules["max_cv"]):
        flags.append("cv>33%")
    return "; ".join(flags)


def weighted_total(df: pd.DataFrame, mask: pd.Series | np.ndarray | None = None) -> float:
    if mask is None:
        return float(df["weight"].sum())
    return float(df.loc[mask, "weight"].sum())


def design_total_cv(base: pd.DataFrame, mask: pd.Series | np.ndarray) -> tuple[float, float | None, float | None]:
    """Approximate Taylor SE/CV for a weighted total using NEISS stratum/PSU variables.

    The calculation aggregates weighted domain totals by stratum and PSU:
    sum_h m_h / (m_h - 1) * sum_k (t_hk - mean_h)^2.
    If design variables are not usable or fewer than two PSUs are observed in all
    relevant strata, SE/CV are returned as missing.
    """

    est = weighted_total(base, mask)
    if "stratum" not in base or "psu" not in base or base["stratum"].isna().all() or base["psu"].isna().all():
        return est, None, None
    tmp = base[["stratum", "psu", "weight"]].copy()
    tmp["_domain_w"] = np.where(np.asarray(mask), tmp["weight"].to_numpy(dtype=float), 0.0)
    grouped = tmp.groupby(["stratum", "psu"], dropna=False)["_domain_w"].sum().reset_index()
    var = 0.0
    usable = False
    for _, g in grouped.groupby("stratum", dropna=False):
        m = len(g)
        if m <= 1:
            continue
        usable = True
        vals = g["_domain_w"].to_numpy(dtype=float)
        var += m / (m - 1) * float(((vals - vals.mean()) ** 2).sum())
    if not usable or est <= 0:
        return est, None, None
    se = math.sqrt(max(var, 0.0))
    return est, se, se / est


def summarize_domain(
    base: pd.DataFrame,
    group_col: str,
    config: dict[str, Any],
    denominator_mask: pd.Series | np.ndarray | None = None,
) -> pd.DataFrame:
    if denominator_mask is None:
        denominator_mask = np.ones(len(base), dtype=bool)
    denom_est = weighted_total(base, denominator_mask)
    rows = []
    values = [v for v in base.loc[denominator_mask, group_col].dropna().unique()]
    for val in sorted(values, key=lambda x: str(x)):
        mask = (base[group_col] == val) & denominator_mask
        n = int(mask.sum())
        est, se, cv = design_total_cv(base, mask)
        pct = est / denom_est if denom_est else np.nan
        rows.append(
            {
                group_col: val,
                "unweighted_n": n,
                "weighted_estimate": est,
                "se": se,
                "cv": cv,
                "weighted_percent": pct,
                "instability_flag": instability_flag(n, est, cv, config),
            }
        )
    return pd.DataFrame(rows)


def safe_read_csv(path: Path, **kwargs: Any) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    try:
        return pd.read_csv(path, **kwargs)
    except pd.errors.EmptyDataError:
        return pd.DataFrame()


def source_note() -> str:
    return (
        "Source: CPSC NEISS annual archived data. Weighted estimates use NEISS "
        "sample weights. Instability flags follow CPSC criteria."
    )


def assert_all_years_present(years: Iterable[int]) -> None:
    missing = [str(annual_raw_path(int(y))) for y in years if not annual_raw_path(int(y)).exists()]
    if missing:
        raise FileNotFoundError(
            "Missing raw NEISS files. Follow data/raw/README_download.md.\n" + "\n".join(missing)
        )


def count_lines(path: Path) -> int:
    with path.open("rb") as f:
        return max(sum(1 for _ in f) - 1, 0)


def write_csv_rows(rows: list[dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
