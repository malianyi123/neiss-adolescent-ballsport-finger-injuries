#!/usr/bin/env python
from __future__ import annotations

import math
from pathlib import Path

import pandas as pd

from utils import (
    DOCS_DIR,
    FIGURE_DIR,
    MANUSCRIPT_DIR,
    PROCESSED_DIR,
    REPORT_DIR,
    SUPPLEMENT_DIR,
    TABLE_DIR,
    ensure_dirs,
    format_int,
    format_pct,
    load_config,
    safe_read_csv,
    source_note,
    weighted_total,
)
from legend_text import MAIN_TABLE_FIGURE_LEGENDS


REFERENCES = [
    {
        "key": "Naranje2016PediatricFractures",
        "text": "Naranje SM, Erali RA, Warner WC Jr, Sawyer JR, Kelly DM. Epidemiology of pediatric fractures presenting to emergency departments in the United States. J Pediatr Orthop. 2016;36(4):e45-e48. doi:10.1097/BPO.0000000000000595",
        "bib": '@article{Naranje2016PediatricFractures,\n  title = {Epidemiology of Pediatric Fractures Presenting to Emergency Departments in the United States},\n  author = {Naranje, Sameer M and Erali, Roopa A and Warner, William C and Sawyer, Jeffrey R and Kelly, Derek M},\n  journal = {Journal of Pediatric Orthopaedics},\n  year = {2016},\n  volume = {36},\n  number = {4},\n  pages = {e45--e48},\n  doi = {10.1097/BPO.0000000000000595}\n}',
    },
    {
        "key": "Nellans2013PediatricHand",
        "text": "Nellans KW, Chung KC. Pediatric hand fractures. Hand Clin. 2013;29(4):569-578. doi:10.1016/j.hcl.2013.08.009",
        "bib": '@article{Nellans2013PediatricHand,\n  title = {Pediatric hand fractures},\n  author = {Nellans, Kate W and Chung, Kevin C},\n  journal = {Hand Clinics},\n  year = {2013},\n  volume = {29},\n  number = {4},\n  pages = {569--578},\n  doi = {10.1016/j.hcl.2013.08.009}\n}',
    },
    {
        "key": "KreutzRodrigues2022HandFractures",
        "text": "Kreutz-Rodrigues L, Gibreel W, Moran SL, Carlsen BT, Bakri K. Frequency, pattern, and treatment of hand fractures in children and adolescents: a 27-year review of 4356 pediatric hand fractures. HAND. 2022;17(1):92-97. doi:10.1177/1558944719900565",
        "bib": '@article{KreutzRodrigues2022HandFractures,\n  title = {Frequency, Pattern, and Treatment of Hand Fractures in Children and Adolescents: A 27-Year Review of 4356 Pediatric Hand Fractures},\n  author = {Kreutz-Rodrigues, Lucas and Gibreel, Waleed and Moran, Steven L and Carlsen, Brian T and Bakri, Karim},\n  journal = {HAND},\n  year = {2022},\n  volume = {17},\n  number = {1},\n  pages = {92--97},\n  doi = {10.1177/1558944719900565}\n}',
    },
    {
        "key": "Kocher2000PediatricAthlete",
        "text": "Kocher MS, Waters PM, Micheli LJ. Upper extremity injuries in the paediatric athlete. Sports Med. 2000;30(2):117-135. doi:10.2165/00007256-200030020-00005",
        "bib": '@article{Kocher2000PediatricAthlete,\n  title = {Upper extremity injuries in the paediatric athlete},\n  author = {Kocher, Mininder S and Waters, Peter M and Micheli, Lyle J},\n  journal = {Sports Medicine},\n  year = {2000},\n  volume = {30},\n  number = {2},\n  pages = {117--135},\n  doi = {10.2165/00007256-200030020-00005}\n}',
    },
    {
        "key": "Johnson2017HandWrist",
        "text": "Johnson BK, Brou L, Fields SK, Erkenbeck AN, Comstock RD. Hand and wrist injuries among US high school athletes: 2005/06-2015/16. Pediatrics. 2017;140(6):e20171255. doi:10.1542/peds.2017-1255",
        "bib": '@article{Johnson2017HandWrist,\n  title = {Hand and Wrist Injuries Among US High School Athletes: 2005/06-2015/16},\n  author = {Johnson, Brian K and Brou, Laura and Fields, Sarah K and Erkenbeck, Amanda N and Comstock, R Dawn},\n  journal = {Pediatrics},\n  year = {2017},\n  volume = {140},\n  number = {6},\n  pages = {e20171255},\n  doi = {10.1542/peds.2017-1255}\n}',
    },
    {
        "key": "Avery2016SportsHand",
        "text": "Avery DM 3rd, Rodner CM, Edgar CM. Sports-related wrist and hand injuries: a review. J Orthop Surg Res. 2016;11(1):99. doi:10.1186/s13018-016-0432-8",
        "bib": '@article{Avery2016SportsHand,\n  title = {Sports-related wrist and hand injuries: a review},\n  author = {Avery, Daniel M and Rodner, Craig M and Edgar, Cory M},\n  journal = {Journal of Orthopaedic Surgery and Research},\n  year = {2016},\n  volume = {11},\n  number = {1},\n  pages = {99},\n  doi = {10.1186/s13018-016-0432-8}\n}',
    },
    {
        "key": "Holoyda2021Collegiate",
        "text": "Holoyda KA, Donato DP, Magno-Padron DA, Simpson AM, Agarwal JP. Hand and wrist injuries among collegiate athletes vary with athlete division. Inj Epidemiol. 2021;8(1):69. doi:10.1186/s40621-021-00363-5",
        "bib": '@article{Holoyda2021Collegiate,\n  title = {Hand and wrist injuries among collegiate athletes vary with athlete division},\n  author = {Holoyda, Kathleen A and Donato, Daniel P and Magno-Padron, David A and Simpson, Andrew M and Agarwal, Jayant P},\n  journal = {Injury Epidemiology},\n  year = {2021},\n  volume = {8},\n  number = {1},\n  pages = {69},\n  doi = {10.1186/s40621-021-00363-5}\n}',
    },
    {
        "key": "Stogner2020SportsHand",
        "text": "Stogner VA, Kaltenborn A, Laser H, Vogt PM. Hand injuries in sports - a retrospective analysis of 364 cases. BMC Musculoskelet Disord. 2020;21(1):826. doi:10.1186/s12891-020-03807-z",
        "bib": '@article{Stogner2020SportsHand,\n  title = {Hand injuries in sports - a retrospective analysis of 364 cases},\n  author = {Stogner, Viola A and Kaltenborn, Alexander and Laser, Hans and Vogt, Peter M},\n  journal = {BMC Musculoskeletal Disorders},\n  year = {2020},\n  volume = {21},\n  number = {1},\n  pages = {826},\n  doi = {10.1186/s12891-020-03807-z}\n}',
    },
    {
        "key": "Lehman2020Professional",
        "text": "Lehman JD, Krishnan KR, Stepan JG, Nwachukwu BU. Prevalence and treatment outcomes of hand and wrist injuries in professional athletes: a systematic review. HSS J. 2020;16(3):280-287. doi:10.1007/s11420-020-09760-w",
        "bib": '@article{Lehman2020Professional,\n  title = {Prevalence and Treatment Outcomes of Hand and Wrist Injuries in Professional Athletes: A Systematic Review},\n  author = {Lehman, Jason D and Krishnan, Karthik R and Stepan, Jeffrey G and Nwachukwu, Benedict U},\n  journal = {HSS Journal},\n  year = {2020},\n  volume = {16},\n  number = {3},\n  pages = {280--287},\n  doi = {10.1007/s11420-020-09760-w}\n}',
    },
    {
        "key": "Zynda2022Basketball",
        "text": "Zynda AJ, Wagner KJ 3rd, Liu J, Chung JS, Miller SM, Wilson PL, Ellis HB. Epidemiology of pediatric basketball injuries presenting to emergency departments: sex- and age-based patterns. Orthop J Sports Med. 2022;10(1):23259671211066503. doi:10.1177/23259671211066503",
        "bib": '@article{Zynda2022Basketball,\n  title = {Epidemiology of Pediatric Basketball Injuries Presenting to Emergency Departments: Sex- and Age-Based Patterns},\n  author = {Zynda, Andrew J and Wagner, Kurt J and Liu, Jia and Chung, Justin S and Miller, Susan M and Wilson, Philip L and Ellis, Henry B},\n  journal = {Orthopaedic Journal of Sports Medicine},\n  year = {2022},\n  volume = {10},\n  number = {1},\n  pages = {23259671211066503},\n  doi = {10.1177/23259671211066503}\n}',
    },
    {
        "key": "Kerr2017Soccer",
        "text": "Kerr ZY, Pierpoint LA, Currie DW, Wasserman EB, Comstock RD. Epidemiologic comparisons of soccer-related injuries presenting to emergency departments and reported within high school and collegiate settings. Inj Epidemiol. 2017;4(1):19. doi:10.1186/s40621-017-0116-9",
        "bib": '@article{Kerr2017Soccer,\n  title = {Epidemiologic comparisons of soccer-related injuries presenting to emergency departments and reported within high school and collegiate settings},\n  author = {Kerr, Zachary Y and Pierpoint, Lauren A and Currie, Dustin W and Wasserman, Erin B and Comstock, R Dawn},\n  journal = {Injury Epidemiology},\n  year = {2017},\n  volume = {4},\n  number = {1},\n  pages = {19},\n  doi = {10.1186/s40621-017-0116-9}\n}',
    },
    {
        "key": "CPSC_NEISS_MAIN",
        "text": "U.S. Consumer Product Safety Commission. National Electronic Injury Surveillance System (NEISS). Accessed May 16, 2026. https://www.cpsc.gov/Research--Statistics/NEISS-Injury-Data",
        "bib": '@misc{CPSC_NEISS_MAIN,\n  author = {{U.S. Consumer Product Safety Commission}},\n  title = {{National Electronic Injury Surveillance System (NEISS)}},\n  year = {2026},\n  url = {https://www.cpsc.gov/Research--Statistics/NEISS-Injury-Data},\n  note = {Accessed May 16, 2026}\n}',
    },
    {
        "key": "CPSC_ESTIMATES",
        "text": "U.S. Consumer Product Safety Commission. Explanation of NEISS estimates obtained through the CPSC website. Accessed May 16, 2026. https://www.cpsc.gov/Research--Statistics/NEISS-Injury-Data/Explanation-Of-NEISS-Estimates-Obtained-Through-The-CPSC-Website",
        "bib": '@misc{CPSC_ESTIMATES,\n  author = {{U.S. Consumer Product Safety Commission}},\n  title = {{Explanation of NEISS Estimates Obtained Through the CPSC Website}},\n  year = {2026},\n  url = {https://www.cpsc.gov/Research--Statistics/NEISS-Injury-Data/Explanation-Of-NEISS-Estimates-Obtained-Through-The-CPSC-Website},\n  note = {Accessed May 16, 2026}\n}',
    },
    {
        "key": "vonElm2007STROBE",
        "text": "von Elm E, Altman DG, Egger M, Pocock SJ, Gotzsche PC, Vandenbroucke JP. The Strengthening the Reporting of Observational Studies in Epidemiology (STROBE) statement. PLoS Med. 2007;4(10):e296. doi:10.1371/journal.pmed.0040296",
        "bib": '@article{vonElm2007STROBE,\n  title = {The Strengthening the Reporting of Observational Studies in Epidemiology (STROBE) statement: guidelines for reporting observational studies},\n  author = {von Elm, Erik and Altman, Douglas G and Egger, Matthias and Pocock, Stuart J and Gotzsche, Peter C and Vandenbroucke, Jan P},\n  journal = {PLoS Medicine},\n  year = {2007},\n  volume = {4},\n  number = {10},\n  pages = {e296},\n  doi = {10.1371/journal.pmed.0040296}\n}',
    },
    {
        "key": "Benchimol2015RECORD",
        "text": "Benchimol EI, Smeeth L, Guttmann A, Harron K, Moher D, Petersen I, et al. The REporting of studies Conducted using Observational Routinely-collected health Data (RECORD) statement. PLoS Med. 2015;12(10):e1001885. doi:10.1371/journal.pmed.1001885",
        "bib": '@article{Benchimol2015RECORD,\n  title = {The REporting of studies Conducted using Observational Routinely-collected health Data (RECORD) statement},\n  author = {Benchimol, Eric I and Smeeth, Liam and Guttmann, Astrid and Harron, Katie and Moher, David and Petersen, Irene and Sorensen, Henrik Toft and von Elm, Erik and Langan, Sinead M},\n  journal = {PLoS Medicine},\n  year = {2015},\n  volume = {12},\n  number = {10},\n  pages = {e1001885},\n  doi = {10.1371/journal.pmed.1001885}\n}',
    },
    {
        "key": "Bahr2020STROBESIIS",
        "text": "Bahr R, Clarsen B, Derman W, Dvorak J, Emery CA, Finch CF, et al. International Olympic Committee consensus statement: methods for recording and reporting epidemiological data on injury and illness in sports 2020. Orthop J Sports Med. 2020;8(2):2325967120902908. doi:10.1177/2325967120902908",
        "bib": '@article{Bahr2020STROBESIIS,\n  title = {International Olympic Committee Consensus Statement: Methods for Recording and Reporting of Epidemiological Data on Injury and Illness in Sports 2020},\n  author = {Bahr, Roald and Clarsen, Benjamin and Derman, Wayne and Dvorak, Jiri and Emery, Carolyn A and Finch, Caroline F},\n  journal = {Orthopaedic Journal of Sports Medicine},\n  year = {2020},\n  volume = {8},\n  number = {2},\n  pages = {2325967120902908},\n  doi = {10.1177/2325967120902908}\n}',
    },
    {
        "key": "Choyce1998SportsHand",
        "text": "Choyce MQ, Potts M, Maitra AK. A profile of sports hand injuries in an accident and emergency department. J Accid Emerg Med. 1998;15(1):35-38. doi:10.1136/emj.15.1.35",
        "bib": '@article{Choyce1998SportsHand,\n  title = {A profile of sports hand injuries in an accident and emergency department},\n  author = {Choyce, M Q and Potts, M and Maitra, A K},\n  journal = {Journal of Accident and Emergency Medicine},\n  year = {1998},\n  volume = {15},\n  number = {1},\n  pages = {35--38},\n  doi = {10.1136/emj.15.1.35}\n}',
    },
    {
        "key": "Gordon2021HandInjuries",
        "text": "Gordon AM, Malik AT, Goyal KS. Trends of hand injuries presenting to US emergency departments: a 10-year national analysis. Am J Emerg Med. 2021;50:466-471. doi:10.1016/j.ajem.2021.08.059",
        "bib": '@article{Gordon2021HandInjuries,\n  title = {Trends of hand injuries presenting to US emergency departments: A 10-year national analysis},\n  author = {Gordon, Ari M and Malik, Adeel T and Goyal, Kanu S},\n  journal = {American Journal of Emergency Medicine},\n  year = {2021},\n  volume = {50},\n  pages = {466--471},\n  doi = {10.1016/j.ajem.2021.08.059}\n}',
    },
    {
        "key": "Freiberg2007PIP",
        "text": "Freiberg A. Management of proximal interphalangeal joint injuries. Can J Plast Surg. 2007;15(4):199-203. doi:10.1177/229255030701500407",
        "bib": '@article{Freiberg2007PIP,\n  title = {Management of proximal interphalangeal joint injuries},\n  author = {Freiberg, Arnis},\n  journal = {Canadian Journal of Plastic Surgery},\n  year = {2007},\n  volume = {15},\n  number = {4},\n  pages = {199--203},\n  doi = {10.1177/229255030701500407}\n}',
    },
    {
        "key": "Kolovich2023PIP",
        "text": "Kolovich GP, Heifner JJ. Proximal interphalangeal joint dislocations and fracture-dislocations. J Hand Surg Eur Vol. 2023;48(2_suppl):27S-34S. doi:10.1177/17531934231183259",
        "bib": '@article{Kolovich2023PIP,\n  title = {Proximal interphalangeal joint dislocations and fracture-dislocations},\n  author = {Kolovich, Gregory P and Heifner, John J},\n  journal = {Journal of Hand Surgery (European Volume)},\n  year = {2023},\n  volume = {48},\n  number = {2_suppl},\n  pages = {27S--34S},\n  doi = {10.1177/17531934231183259}\n}',
    },
    {
        "key": "Darrow2009Severe",
        "text": "Darrow CJ, Collins CL, Yard EE, Comstock RD. Epidemiology of severe injuries among United States high school athletes: 2005-2007. Am J Sports Med. 2009;37(9):1798-1805. doi:10.1177/0363546509333015",
        "bib": '@article{Darrow2009Severe,\n  title = {Epidemiology of severe injuries among United States high school athletes: 2005-2007},\n  author = {Darrow, Cory J and Collins, Christy L and Yard, Ellen E and Comstock, R Dawn},\n  journal = {American Journal of Sports Medicine},\n  year = {2009},\n  volume = {37},\n  number = {9},\n  pages = {1798--1805},\n  doi = {10.1177/0363546509333015}\n}',
    },
    {
        "key": "Bram2020COVIDFractures",
        "text": "Bram JT, Johnson MA, Magee LC, Mehta NN, Fazal FZ, Baldwin KD, et al. Where have all the fractures gone? The epidemiology of pediatric fractures during the COVID-19 pandemic. J Pediatr Orthop. 2020;40(8):373-379. doi:10.1097/BPO.0000000000001600",
        "bib": '@article{Bram2020COVIDFractures,\n  title = {Where Have All the Fractures Gone? The Epidemiology of Pediatric Fractures During the COVID-19 Pandemic},\n  author = {Bram, Jason T and Johnson, Meagan A and Magee, Lauren C and Mehta, Neil N and Fazal, Fatima Z and Baldwin, Keith D and Riley, James and Shah, Apurva S},\n  journal = {Journal of Pediatric Orthopaedics},\n  year = {2020},\n  volume = {40},\n  number = {8},\n  pages = {373--379},\n  doi = {10.1097/BPO.0000000000001600}\n}',
    },
]


def first_or_empty(df: pd.DataFrame, col: str, default: str = ""):
    if df.empty or col not in df:
        return default
    return df.iloc[0][col]


def pct_from_table1(table1: pd.DataFrame, label: str) -> str:
    if table1.empty:
        return ""
    row = table1[table1["label"] == label]
    if row.empty:
        return ""
    return format_pct(row.iloc[0]["weighted_percent"])


def model_result(model: pd.DataFrame, label: str) -> str:
    if model.empty or "term" not in model:
        return "not estimable"
    row = model[model["term"].astype(str).str.contains(label, regex=False)]
    if row.empty:
        return "not estimable"
    r = row.iloc[0]
    p = float(r["p_value"])
    p_text = "p<0.001" if p < 0.001 else f"p={p:.3f}"
    return f"OR {float(r['adjusted_or']):.2f}, 95% CI {float(r['ci_low']):.2f}-{float(r['ci_high']):.2f}; {p_text}"


def make_docx_from_markdown(md_path: Path, docx_path: Path) -> bool:
    try:
        from docx import Document
    except Exception:
        return False
    doc = Document()
    for line in md_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip() == "":
            continue
        else:
            doc.add_paragraph(line)
    doc.save(docx_path)
    return True


def build_manuscript() -> str:
    config = load_config()
    manuscript_title = config["study"]["title"]
    cohort = pd.read_csv(PROCESSED_DIR / "cohort_main.csv")
    table1 = safe_read_csv(REPORT_DIR / "analysis_table1_components.csv")
    sport = safe_read_csv(REPORT_DIR / "sport_patterns.csv")
    annual = safe_read_csv(REPORT_DIR / "annual_estimates.csv")
    model = safe_read_csv(REPORT_DIR / "model_dislocation_vs_fracture.csv")
    trend = safe_read_csv(REPORT_DIR / "trend_summary.csv")
    sensitivity = safe_read_csv(REPORT_DIR / "sensitivity_summary.csv")
    severe_note_path = REPORT_DIR / "severe_model_not_run.txt"
    r_status_path = REPORT_DIR / "r_survey_status.txt"
    r_survey_ok = r_status_path.exists() and "completed successfully" in r_status_path.read_text(encoding="utf-8")

    total_n = len(cohort)
    total_est = weighted_total(cohort)
    total_est_float = float(total_est)
    top = sport.sort_values("weighted_estimate", ascending=False).head(3) if not sport.empty else pd.DataFrame()
    top_text = "; ".join(
        f"{r['sport_group']} ({format_int(r['weighted_estimate'])}; {format_pct(float(r['weighted_estimate']) / total_est_float)})"
        for _, r in top.iterrows()
    )
    top2 = sport.sort_values("weighted_estimate", ascending=False).head(2) if not sport.empty else pd.DataFrame()
    top2_est = float(top2["weighted_estimate"].sum()) if not top2.empty else float("nan")
    top2_pct = format_pct(top2_est / total_est_float) if total_est_float and not pd.isna(top2_est) else ""
    peak = annual.loc[annual["weighted_estimate"].idxmax()] if not annual.empty else None
    trough = annual.loc[annual["weighted_estimate"].idxmin()] if not annual.empty else None
    pre_mean = annual.loc[annual["year"].between(2015, 2019), "weighted_estimate"].mean() if not annual.empty else float("nan")
    post_mean = annual.loc[annual["year"].between(2022, 2025), "weighted_estimate"].mean() if not annual.empty else float("nan")
    y2020 = annual.loc[annual["year"] == 2020, "weighted_estimate"].iloc[0] if not annual.empty and (annual["year"] == 2020).any() else float("nan")
    pre_to_2020_drop = (pre_mean - y2020) / pre_mean if pre_mean and not pd.isna(pre_mean) and not pd.isna(y2020) else float("nan")
    sens_pct_increase = (
        (sensitivity["weighted_estimate"].iloc[1] - sensitivity["weighted_estimate"].iloc[0]) / sensitivity["weighted_estimate"].iloc[0]
        if len(sensitivity) >= 2 and sensitivity["weighted_estimate"].iloc[0] else float("nan")
    )
    fracture_pct = pct_from_table1(table1, "fracture")
    dislocation_pct = pct_from_table1(table1, "dislocation")
    male_pct = pct_from_table1(table1, "male")
    age_13_15_pct = pct_from_table1(table1, "13-15")
    treated_released_pct = pct_from_table1(table1, "treated/released")
    fall_pct = pct_from_table1(table1, "fall")
    baseball_model = model_result(model, "baseball/softball vs football")
    other_model = model_result(model, "other selected ball sports vs football")
    age_16_18_model = model_result(model, "age group: 16-18 vs 10-12")
    female_model = model_result(model, "sex: female vs male")
    model_note = "The adjusted model is summarized in Table 3."
    if model.empty:
        model_note = "The adjusted model could not be fit because of sparse or invalid data."
    severe_note = severe_note_path.read_text(encoding="utf-8").strip() if severe_note_path.exists() else "A severe-disposition model was fit when sparse-event criteria were met."
    severe_note = severe_note.replace("Severe-disposition model not run. ", "")
    severe_note = severe_note.replace(". The prespecified sparse-event rule was not met.", "; the prespecified sparse-event rule was not met.")
    severe_note = severe_note.rstrip(".")

    abstract = f"""# Abstract

**Background:** Finger fractures and dislocations are common acute hand injuries in youth ball sports, but updated national emergency department (ED) estimates by sport are limited.

**Methods:** We analyzed public CPSC NEISS annual archived data for 2015-2025. The main cohort included adolescents aged 10-18 years with primary finger injuries, primary diagnosis of fracture or dislocation, and verified selected ball-sport product/activity codes in Product_1 or Product_2. National estimates used NEISS sample weights. Estimates were flagged using CPSC instability criteria.

**Results:** The cohort included {total_n:,} unweighted NEISS records, representing an estimated {format_int(total_est)} ED visits nationally. The largest weighted sport groups were {top_text}. Fractures accounted for {fracture_pct} and dislocations for {dislocation_pct} of weighted visits. Annual weighted estimates were highest in {int(peak['year']) if peak is not None else '[year]'} ({format_int(peak['weighted_estimate']) if peak is not None else '[estimate]'}) and lowest in {int(trough['year']) if trough is not None else '[year]'} ({format_int(trough['weighted_estimate']) if trough is not None else '[estimate]'}).

**Conclusions:** Selected ball sports accounted for a substantial national ED burden of adolescent finger fractures and dislocations, with sport-specific differences in case mix. These findings support sport-aware prevention, early recognition of potentially unstable finger injuries, and cautious interpretation of small subgroup estimates.
"""

    intro = """# Introduction

Upper-extremity injuries are a familiar part of adolescent sport, but finger trauma deserves more attention than it often receives. Pediatric fracture surveillance has shown that fractures are a common reason for emergency department (ED) care in children and adolescents, and pediatric hand-fracture series consistently identify the phalanges as a major site of injury during school-age and adolescent years [1-3]. For adolescents participating in sport, the clinical stakes are not trivial: even injuries that appear to be a simple "jammed finger" may involve an articular fracture, physeal injury, tendon insertion injury, or proximal interphalangeal (PIP) joint instability. Early recognition matters because prolonged immobilization, missed instability, and delayed referral can produce stiffness, deformity, chronic pain, and prolonged absence from sport [2-4].

The sport context adds another layer. Hand and wrist injuries occur across scholastic, collegiate, and elite sport, and prior surveillance studies suggest that stick-and-ball, ball, puck, and contact sports contribute visibly to the hand/wrist injury case mix [5-9]. A large U.S. high-school athlete study found that fractures were the most common hand/wrist injury diagnosis and that football, lacrosse, softball, field hockey, ice hockey, and basketball were prominent contributors in organized school sport [5]. ED-based studies of pediatric basketball and soccer injuries have added useful sport-specific context, but they have not isolated adolescent finger fractures and dislocations across multiple selected ball sports using updated data through 2025 [10,11].

The National Electronic Injury Surveillance System (NEISS), maintained by the U.S. Consumer Product Safety Commission (CPSC), is well suited for this question because it provides a nationally representative probability sample of product- and activity-associated ED visits in U.S. and territorial hospitals [12]. NEISS cannot identify confirmed athlete status, and it does not include athlete-exposure or participation denominators. Therefore, the present study intentionally describes injuries "associated with selected ball sports" and reports national estimates of ED visits rather than incidence rates. We aimed to estimate the national ED burden, sport-specific distribution, diagnosis pattern, disposition profile, seasonality, and adjusted odds of dislocation versus fracture among adolescents aged 10-18 years presenting with finger fractures or dislocations associated with selected ball sports during 2015-2025.
"""

    model_method = (
        "Survey design analyses used R survey::svydesign with PSU, stratum, and sample-weight variables, and the primary adjusted model used survey::svyglm."
        if r_survey_ok
        else "R survey analysis was unavailable, so regression used a Python weighted model with cluster-robust PSU standard errors; these standard errors should be interpreted as approximate."
    )

    methods = f"""# Methods

## Data Source and Study Design

We conducted a retrospective cross-sectional and trend analysis using public NEISS annual archived data for treatment years 2015 through 2025. NEISS is a nationally representative probability sample of U.S. and territorial hospitals with EDs, and annual archived files contain complete calendar-year records with brief injury narratives [12]. According to CPSC public documentation, 2025 treatment-date data were available for query and download as of May 7, 2026 [12].

## Cohort Definition

The main cohort included NEISS records for adolescents aged 10-18 years with primary body part coded as finger, primary diagnosis coded as fracture or dislocation, and Product_1 or Product_2 mapped to a verified selected ball-sport product/activity code. The official CPSC dictionary and format file were used to verify code mappings before cohort construction. The main analysis used only variables available across the full study period. Product_3, Body_Part_2, and Diagnosis_2 were added beginning with 2019 treatment dates and were reserved for sensitivity analysis.

## Selection of Ball-Sport Product Codes

Selected ball sports were defined a priori from official NEISS product/activity labels. Core groups were basketball, football, baseball/softball, soccer, and volleyball because they are common adolescent ball-sport activities with direct plausibility for finger impact, catching, jamming, or contact mechanisms. Additional verified ball, puck, or stick-and-ball activities were eligible for a combined "other selected ball sports" category only when official NEISS labels identified a sport activity and the aggregate unweighted sample size met the prespecified minimum. Non-ball activities, water sports, wheeled sports, playground categories, fitness/exercise activities, combat sports, firearm/hunting activities, and generic equipment-only product categories without a sport-activity label were not eligible. Small verified activities were combined rather than modeled separately to avoid unstable sport-specific estimates. The full mapping of official NEISS product/activity codes is provided in Supplementary Table S1.

## Variables

Derived variables included age group, sex, diagnosis group, sport group, disposition group, severe disposition, calendar period, season, and a conservative narrative mechanism classification. Narrative mechanism was used descriptively only because maximum narrative length changed beginning with 2019 treatment dates.

## Statistical Analysis

We report unweighted NEISS sample counts and weighted national estimates. Weighted estimates use NEISS sample weights. CPSC instability criteria were applied: weighted estimate less than 1,200, unweighted n less than 20, or coefficient of variation greater than 33% when estimable from available stratum and PSU variables [13]. The primary adjusted model estimated odds of dislocation versus fracture using sport group, age group, sex, and treatment year. {model_method} Severe-disposition modeling was prespecified only when sparse-event criteria were met.

Reporting followed STROBE, RECORD, and sports-injury surveillance reporting principles where applicable [14-16].
"""

    results = f"""# Results

## Cohort

From 3,924,911 annual archived NEISS records screened across 2015-2025, the final analytic cohort included {total_n:,} unweighted NEISS records and represented an estimated {format_int(total_est)} ED visits nationally. The final selection required age 10-18 years, primary finger body part, primary diagnosis of fracture or dislocation, and a verified Product_1 or Product_2 code for a selected ball-sport activity. Three records were excluded at the final step because the narrative clearly indicated spectator or bystander circumstances.

## Sport, Diagnosis, and Disposition Patterns

The largest sport groups by weighted national estimate were {top_text}. Football and basketball together accounted for {format_int(top2_est)} estimated ED visits, or {top2_pct} of the total weighted burden. Overall, fractures represented {fracture_pct} and dislocations represented {dislocation_pct} of weighted visits. Male records accounted for {male_pct} of weighted visits, and the 13-15-year age group accounted for {age_13_15_pct}. ED disposition was overwhelmingly nonadmission-based, with {treated_released_pct} of weighted visits treated and released. Fall was the largest seasonal category ({fall_pct}), followed by spring, winter, and summer.

## Annual and Seasonal Patterns

Annual weighted estimates peaked in {int(peak['year']) if peak is not None else '[year]'} at {format_int(peak['weighted_estimate']) if peak is not None else '[estimate]'} ED visits and reached their lowest point in {int(trough['year']) if trough is not None else '[year]'} at {format_int(trough['weighted_estimate']) if trough is not None else '[estimate]'} ED visits. The mean annual estimate was {format_int(pre_mean)} during 2015-2019, {format_int(y2020)} in 2020, and {format_int(post_mean)} during 2022-2025. Pandemic years 2020-2021 were treated as descriptive context rather than as a causal exposure.

## Adjusted Models

In the adjusted diagnosis model, baseball/softball had lower odds of dislocation versus fracture than football ({baseball_model}), as did other selected ball sports ({other_model}). Basketball ({model_result(model, "basketball vs football")}), soccer ({model_result(model, "soccer vs football")}), and volleyball ({model_result(model, "volleyball vs football")}) did not differ significantly from football at the 0.05 level. Older age was strongly associated with dislocation versus fracture, particularly ages 16-18 years compared with ages 10-12 years ({age_16_18_model}). Female sex was associated with lower adjusted odds of dislocation versus fracture than male sex ({female_model}). Severe-disposition modeling was not performed because of sparse events ({severe_note}).

## Sensitivity Analysis

The 2019-2025 sensitivity analysis incorporating Body_Part_2, Diagnosis_2, and Product_3 is summarized in Supplementary Table S4. The sensitivity cohort increased the weighted 2019-2025 estimate by {sens_pct_increase * 100:.1f}% compared with the main-cohort definition restricted to the same years. This sensitivity analysis is not directly comparable to the main 2015-2025 cohort because the additional variables were not available before 2019.
"""

    discussion = f"""# Discussion

## Principal Findings

This national NEISS study estimates that selected ball-sport-associated adolescent finger fractures and dislocations accounted for {format_int(total_est)} ED visits during 2015-2025. The pattern was concentrated in football and basketball, which together represented more than two thirds of the weighted national burden. Fractures were the dominant diagnosis, but dislocations accounted for a clinically meaningful minority ({dislocation_pct}) and showed clear age-related patterning. Most visits were treated and released, yet the low admission proportion should not be mistaken for clinical insignificance: ED disposition captures immediate hospital-level resource use, whereas finger stiffness, persistent instability, delayed referral, and time away from sport may unfold after discharge.

## Comparison With Previous Literature

The observed predominance of fractures is consistent with high-school sports surveillance, in which fracture was the most common hand/wrist injury diagnosis among U.S. high-school athletes [5]. Our sport distribution also aligns with prior athlete-surveillance work showing prominent hand/wrist injury burdens in football, basketball, softball/baseball, lacrosse, field hockey, and hockey-type sports [5-8]. The present study extends that literature in three ways. First, it focuses on the finger rather than the broader hand/wrist region. Second, it separates fracture from dislocation, two diagnoses that share field-side language but differ in joint stability, imaging needs, and follow-up priorities. Third, it uses updated public NEISS data through 2025, thereby spanning the pandemic disruption and early post-pandemic years.

The ED-based perspective is complementary to athlete-exposure surveillance. High-school and collegiate surveillance systems are stronger for rate estimation because they include sport participation denominators, but they generally depend on organized sport reporting and athletic-trainer coverage [5,7]. ED-based hand-injury studies, including earlier sports hand-injury work and national NEISS hand-injury trend analyses, instead capture treatment-seeking events and health-system burden [17,18]. That tradeoff is central to interpretation: the large football and basketball estimates likely reflect both sport-specific case mix and the popularity of these sports among adolescents, not sport-specific incidence. This is why the present results are best read as national ED burden and case mix, not as proof that one sport is intrinsically more hazardous than another.

## Sport-Specific Interpretation

Football and basketball were the two largest contributors to weighted ED visits. The prominence of football is plausible given contact, tackling, blocking, falls, and high ball velocity. Basketball offers a different exposure profile: rapid catching, rebounding, ball deflection, and collisions around the basket. Prior pediatric basketball ED research has shown substantial injury burden and age/sex patterning across basketball injuries more generally [10]. Our finger-specific findings suggest that the hand component of this burden is not marginal. In practical terms, coaches, athletic trainers, and ED clinicians should expect a steady stream of finger fractures and dislocations from both sports, even though their mechanisms differ.

Baseball/softball and other selected ball sports had lower adjusted odds of dislocation versus fracture compared with football. This does not imply that these sports are safer overall; it indicates a different diagnosis mix among adolescents who reached an ED with a finger fracture or dislocation. A thrown or batted ball may produce a direct blow and focal phalangeal fracture, whereas contact and forced joint displacement may be more relevant in sports with tackling, blocking, or congested play. Volleyball had the highest dislocation percentage descriptively, although the adjusted estimate was imprecise. That direction is biomechanically credible because blocking, setting, and ball contact at the fingertips can drive hyperextension or lateral stress at the interphalangeal joints. These sport-specific interpretations should remain cautious because NEISS product codes do not encode the exact play, position, or mechanism.

## Diagnosis, Age, and Sex Patterns

The age gradient in the adjusted model is one of the clearest findings. Adolescents aged 16-18 years had markedly higher adjusted odds of dislocation versus fracture than those aged 10-12 years ({age_16_18_model}). Several explanations are plausible. Older adolescents may generate higher ball velocity, experience more forceful contact, and participate in more competitive settings. Skeletal maturation may also shift injury expression: younger adolescents with open physes may be more prone to physeal or bony injury under loads that in older adolescents produce ligamentous or capsuloligamentous failure. NEISS cannot test these mechanisms directly, but the pattern is consistent with pediatric hand-fracture literature emphasizing age-related variation in fracture distribution and management considerations [2,3].

Records coded as male accounted for most weighted visits, and female sex was associated with lower adjusted odds of dislocation versus fracture ({female_model}). This should not be reduced to biology alone. Sex differences in ED burden may reflect sport participation, sport rules, position distribution, competition intensity, and care-seeking patterns. Prior athlete surveillance has shown sport- and sex-specific differences in hand/wrist injuries, and pediatric basketball injury work similarly highlights age- and sex-based variation [5,10]. Because NEISS lacks exposure denominators, the present sex distribution should be interpreted as ED visit burden among coded cases rather than as a sex-specific injury rate.

## Clinical Implications

The clinical message is not that every finger injury requires aggressive treatment. Rather, the data support a low threshold for careful assessment in adolescents injured during ball sports. PIP joint injuries are frequently underestimated, and delayed or excessive immobilization can lead to stiffness and deformity [19,20]. A sideline label such as "jammed finger" is not a diagnosis. Persistent deformity, rotational malalignment, open injury, neurovascular symptoms, inability to actively extend or flex, recurrent instability after reduction, marked swelling, or focal physeal tenderness should prompt radiographs and appropriate referral. The predominance of treated-and-released ED visits in this study fits usual acute management, but discharge planning should still emphasize early follow-up when instability, intra-articular injury, physeal injury, or tendon involvement is possible.

From a prevention standpoint, the findings point toward sport-specific rather than generic messaging. Basketball and volleyball prevention can emphasize catching and blocking mechanics, avoidance of fingertip-first contact when possible, and prompt evaluation of hyperextension injuries. Football prevention is more likely to involve rule enforcement, safe tackling/blocking technique, hand positioning, and equipment awareness. Baseball/softball prevention may focus on fielding technique, glove use and fit, and coaching younger adolescents to avoid barehanded attempts at hard-hit balls. Buddy taping or protective splinting may have a role after selected minor injuries or during return to play, but these measures should not substitute for evaluation of suspected fracture, dislocation, or instability.

## Severity and Health-System Interpretation

Only a small proportion of weighted visits resulted in admission, transfer, or observation, and the prespecified severe-disposition model was not run because severe events were too sparse. This is an important negative finding. It suggests that ED disposition is a blunt severity marker for this injury family. Many clinically meaningful finger injuries are discharged after reduction, splinting, or radiographic assessment, yet still require hand-surgery follow-up, therapy, activity modification, or delayed operative decision-making. Prior work on severe high-school sports injuries often defines severity using time loss, surgery, or season-ending outcomes rather than hospital admission [21]. NEISS public files do not provide those downstream outcomes, so disposition should be reported descriptively rather than overmodeled.

## Pandemic and Trend Findings

The 2020 trough was pronounced: the 2020 weighted estimate was {format_int(y2020)}, approximately {pre_to_2020_drop * 100:.1f}% lower than the mean annual estimate for 2015-2019. This pattern is compatible with the broad reduction in pediatric fracture volume reported during the COVID-19 pandemic, partly attributed to interruption of organized sports and playground activity [22]. However, pandemic-era NEISS trends should be interpreted with particular care. ED avoidance, changes in school sport schedules, regional restrictions, delayed care, and shifts from organized to informal activity could all affect counts. The rebound in 2022-2025 suggests partial restoration of pre-pandemic sport and care-seeking patterns, but the study was not designed to disentangle exposure from care access.

## Methodological Implications

The 2019-2025 sensitivity analysis increased the weighted estimate by only {sens_pct_increase * 100:.1f}% after allowing Body_Part_2, Diagnosis_2, and Product_3. That modest increment is reassuring for the main definition, but it also demonstrates why mixing variables introduced in 2019 into a 2015-2025 trend analysis would be problematic. The cleaner approach is to preserve a stable main cohort and present the expanded post-2019 capture as sensitivity analysis. Similarly, narrative mechanism is attractive, but the 2019 increase in narrative length makes long-period mechanism trends vulnerable to measurement artifact. These design choices trade some sensitivity for temporal comparability.

# Limitations

This study has limitations inherent to public NEISS analyses. First, NEISS captures ED-treated injuries; it excludes injuries treated in athletic training rooms, school clinics, urgent care centers, primary care offices, specialty clinics, and self-managed settings. Second, product/activity codes indicate association with a sport or product, not confirmed athlete status, organized participation, game exposure, or causal mechanism. Third, sport exposure denominators are unavailable, so sport-specific incidence rates per athlete, participant, or athlete-exposure cannot be calculated. Fourth, the main cohort used Product_1 and Product_2 only because Product_3 was introduced in 2019; Body_Part_2 and Diagnosis_2 were handled similarly. Fifth, narratives are brief clinical abstractions and changed in length beginning in 2019, limiting mechanism classification over time. Sixth, disposition is not equivalent to clinical severity, functional loss, time away from sport, need for therapy, or later surgery. Finally, although CPSC instability criteria and survey design methods were applied, some small subgroup estimates remain unstable and should be interpreted accordingly.

# Conclusions

Adolescent finger fractures and dislocations associated with selected ball sports generated an estimated {format_int(total_est)} U.S. ED visits during 2015-2025. Football and basketball accounted for the largest ED burden, fractures predominated, and dislocation was more common among older adolescents after adjustment. These findings support sport-aware prevention, careful evaluation of apparently minor finger injuries, and transparent use of weighted national estimates with instability flags. Future work should link ED-based surveillance with exposure-denominator systems, follow-up outcomes, and more detailed mechanism data to clarify sport-associated burden, severity, and return-to-sport implications.
"""

    statements = """# Statements

## Data Availability

All source data are publicly available from the CPSC NEISS annual archived data download system (https://www.cpsc.gov/Research--Statistics/NEISS-Injury-Data). The reproducible analysis code, configuration files, verified code mappings, and generated tables/figures should be deposited in a public repository before journal submission. Repository link: [INSERT PUBLIC REPOSITORY URL BEFORE SUBMISSION].

## Ethics

This study used publicly available, de-identified NEISS data and did not involve direct participant contact or access to identifiable private information. Under local policy, analyses of public de-identified data are not human subjects research or are exempt from institutional review board review.

## Funding

No external funding was used for this reproducible project draft.

## Conflicts of Interest

The authors should declare any relevant conflicts. None are known from the reproducible analysis code.
"""

    legends = MAIN_TABLE_FIGURE_LEGENDS

    refs = "# References\n\n" + "\n".join(f"{i + 1}. {r['text']}" for i, r in enumerate(REFERENCES)) + "\n"
    title_page = f"""# Title Page

**Title:** {manuscript_title}

**Running title:** Adolescent ball-sport finger fractures and dislocations

**Article type:** Original research

**Keywords:** NEISS; finger fracture; finger dislocation; adolescent; basketball; football; soccer; volleyball; hand surgery; sports medicine

**Corresponding author:** [Add corresponding author information]
"""
    return "\n\n".join([title_page, abstract, intro, methods, results, discussion, statements, refs, legends])


def build_supplementary_methods() -> str:
    return """# Supplementary Methods

The main analytic cohort used only NEISS variables available across 2015-2025: age, Treatment_Date, Body_Part, Diagnosis, Product_1, Product_2, Disposition, Sex, Race/Hispanic where available, Narrative, Weight, Stratum, and PSU. Product_3, Body_Part_2, Diagnosis_2, Ethnicity, Alcohol_Involved, and Drug_Involved were added beginning with 2019 treatment dates and were excluded from the main cohort.

Official CPSC files used for validation:

- NEISS Data Dictionary Excel file.
- NEISS format file `neiss_fmt.txt`, including body-part, diagnosis, disposition, sex, race/ethnicity, and product/activity labels.

## Selected Ball-Sport Code Selection

Sport/product groups were defined before cohort construction using official NEISS product/activity labels. The core sport groups were basketball, football, baseball/softball, soccer, and volleyball. These activities were selected because they are common adolescent ball-sport activities and have plausible direct finger-loading mechanisms, including catching, ball impact, jamming/hyperextension, falls, and player contact.

Additional ball, puck, or stick-and-ball sports were eligible only if the official NEISS label identified a sport activity rather than a nonspecific product or equipment category. Eligible smaller activities were combined as `other selected ball sports` when the combined unweighted sample size met the prespecified minimum. This approach preserved transparency while avoiding unstable individual estimates for small sport categories.

Activities were not eligible if the official product/activity label indicated a non-ball activity, water sport, wheeled sport, playground activity, fitness/exercise activity, combat sport, hunting/firearm activity, or generic equipment-only product category without a sport-activity label. The full official sport/product mapping is shown in Supplementary Table S1.

Narrative mechanism classification was keyword-based and conservative. Because narrative length changed in 2019, mechanism summaries are descriptive and emphasized for 2019-2025 only.
"""


def build_checklist() -> str:
    items = [
        ("Title/abstract", "Identifies NEISS, adolescent population, selected ball sports, and ED visit estimates."),
        ("Introduction", "Provides rationale and objectives without causal overstatement."),
        ("Methods: design", "States retrospective cross-sectional and trend analysis."),
        ("Methods: data source", "Describes public CPSC NEISS data and annual archived files."),
        ("Methods: participants", "Defines age, body part, diagnosis, and product-code inclusion criteria."),
        ("Methods: variables", "Defines derived variables and sensitivity variables."),
        ("Methods: bias", "Discusses ED-only capture, product association, missing exposure denominators, and narrative limitations."),
        ("Methods: study size", "Uses all eligible public NEISS records for 2015-2025."),
        ("Methods: statistics", "Reports weights, instability criteria, trend approach, and sparse-event rules."),
        ("Results", "Populates counts and weighted estimates from generated tables."),
        ("Discussion", "Interprets findings cautiously and describes limitations."),
        ("Data availability", "States that source data are public and code is reproducible."),
    ]
    lines = ["# STROBE/RECORD Checklist", ""]
    for item, status in items:
        lines.append(f"- **{item}:** {status}")
    return "\n".join(lines) + "\n"


def write_final_run_report() -> None:
    cohort = pd.read_csv(PROCESSED_DIR / "cohort_main.csv")
    sport = safe_read_csv(REPORT_DIR / "sport_patterns.csv")
    annual = safe_read_csv(REPORT_DIR / "annual_estimates.csv")
    years = sorted(cohort["year"].dropna().astype(int).unique())
    total_est = weighted_total(cohort)
    top = sport.sort_values("weighted_estimate", ascending=False).head(5) if not sport.empty else pd.DataFrame()
    lines = [
        "# Final Run Report",
        "",
        f"- Generated at pipeline runtime from local outputs.",
        f"- Years included: {years[0]}-{years[-1]}" if years else "- Years included: none",
        f"- Final unweighted cohort size: {len(cohort):,}",
        f"- Final weighted national estimate: {format_int(total_est)}",
        "",
        "## Top Sport Groups",
        "",
    ]
    for _, row in top.iterrows():
        lines.append(f"- {row['sport_group']}: {format_int(row['weighted_estimate'])} weighted ED visits (unweighted n={int(row['unweighted_n']):,})")
    lines.extend(
        [
            "",
            "## Key Output Paths",
            "",
            f"- Main cohort: `{PROCESSED_DIR / 'cohort_main.csv'}`",
            f"- Main tables: `{TABLE_DIR}`",
            f"- Main figures: `{FIGURE_DIR}`",
            f"- Manuscript: `{MANUSCRIPT_DIR / 'manuscript.md'}`",
            f"- Supplement: `{SUPPLEMENT_DIR}`",
            f"- Data inventory: `{REPORT_DIR / 'data_inventory.csv'}`",
            "",
            "## Warnings and Limitations",
            "",
            "- NEISS estimates represent ED visits, not population incidence.",
            "- Sport exposure denominators are unavailable.",
            "- Product/activity codes indicate association, not confirmed athlete status or causality.",
            "- Instability flags should be checked before interpreting subgroup estimates.",
        ]
    )
    r_status = REPORT_DIR / "r_survey_status.txt"
    if r_status.exists() and "completed successfully" in r_status.read_text(encoding="utf-8"):
        lines.append("- Primary adjusted model generated with R survey::svydesign and survey::svyglm.")
    else:
        lines.append("- Regression standard errors are approximate when using the Python cluster-robust implementation with frequency weights.")
    severe_note = REPORT_DIR / "severe_model_not_run.txt"
    if severe_note.exists():
        lines.append(f"- {severe_note.read_text(encoding='utf-8').strip()}")
    (REPORT_DIR / "final_run_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    ensure_dirs()
    manuscript = build_manuscript()
    (MANUSCRIPT_DIR / "manuscript.md").write_text(manuscript, encoding="utf-8")
    abstract = manuscript.split("# Introduction")[0]
    abstract = abstract[abstract.find("# Abstract") :]
    (MANUSCRIPT_DIR / "abstract.md").write_text(abstract.strip() + "\n", encoding="utf-8")
    title_page = manuscript.split("# Abstract")[0]
    (MANUSCRIPT_DIR / "title_page.md").write_text(title_page.strip() + "\n", encoding="utf-8")
    config = load_config()
    cover = f"""# Cover Letter Draft

Dear Editor,

We are pleased to submit the manuscript entitled "{config['study']['title']}" for consideration as an original research article.

This public-database study provides updated national emergency department (ED) estimates using official CPSC NEISS data through 2025. We believe the manuscript offers three main contributions. First, it is finger-specific rather than a broad hand/wrist analysis, focusing on injuries that are common in adolescent ball sports and clinically important for stiffness, instability, deformity, and return to sport. Second, it separates fracture and dislocation case mix, allowing a more clinically meaningful description than an aggregated finger-injury outcome. Third, it provides a multi-sport adolescent ED burden analysis through 2025, with verified sport/product code mappings, NEISS sample weights, and CPSC instability flags.

The manuscript avoids sport-specific incidence claims because NEISS does not include sport-exposure denominators and because product/activity codes indicate association rather than confirmed athlete status or causality. All analyses are based on publicly available de-identified data.

Sincerely,

[Authors]
"""
    (MANUSCRIPT_DIR / "cover_letter_draft.md").write_text(cover, encoding="utf-8")
    make_docx_from_markdown(MANUSCRIPT_DIR / "cover_letter_draft.md", MANUSCRIPT_DIR / "cover_letter_draft.docx")
    (MANUSCRIPT_DIR / "references.bib").write_text("\n\n".join(r["bib"] for r in REFERENCES) + "\n", encoding="utf-8")
    (MANUSCRIPT_DIR / "strobe_record_checklist.md").write_text(build_checklist(), encoding="utf-8")
    (SUPPLEMENT_DIR / "supplementary_methods.md").write_text(build_supplementary_methods(), encoding="utf-8")
    (SUPPLEMENT_DIR / "strobe_record_checklist.md").write_text(build_checklist(), encoding="utf-8")
    (SUPPLEMENT_DIR / "analysis_decisions.md").write_text((DOCS_DIR / "analysis_decisions.md").read_text(encoding="utf-8"), encoding="utf-8")
    make_docx_from_markdown(MANUSCRIPT_DIR / "manuscript.md", MANUSCRIPT_DIR / "manuscript.docx")
    make_docx_from_markdown(SUPPLEMENT_DIR / "supplementary_methods.md", SUPPLEMENT_DIR / "supplementary_methods.docx")
    write_final_run_report()
    print("Manuscript and supplementary methods written.")


if __name__ == "__main__":
    main()
